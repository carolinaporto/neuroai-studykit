"""PPTX parser: one ParsedBlock per slide, slide text + speaker notes.

Locator is {"slide": n, "has_notes": bool} (1-indexed).
"""

import io

from pptx import Presentation

from ..models import Locator, ParsedBlock


def parse_pptx(data: bytes) -> list[ParsedBlock]:
    presentation = Presentation(io.BytesIO(data))
    blocks: list[ParsedBlock] = []

    for index, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        parts = list(texts)
        if notes:
            parts.append(f"Notes: {notes}")
        text = "\n\n".join(parts).strip()
        if not text:
            continue

        blocks.append(ParsedBlock(text=text, locator=Locator(slide=index, has_notes=bool(notes))))

    return blocks
