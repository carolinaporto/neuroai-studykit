"""Generates the synthetic fixture corpus used by ingest tests.

CLAUDE.md #6: content/ (real course material) is gitignored and never enters the repo.
These texts are ours, written for this project, about basic neuroscience/ML concepts, not
GENED 1201 material. The generated files are committed so tests don't depend on rerunning
this script.

Run with: uv run python tests/fixtures/generate_fixtures.py
"""

from pathlib import Path

import pymupdf
from pptx import Presentation

FIXTURES_DIR = Path(__file__).parent

PDF_PAGES = [
    (
        "Neurons and Membrane Potential",
        """A neuron is the basic signaling unit of the nervous system. Like other cells, it
maintains a difference in electrical charge between the inside and the outside of its
membrane. At rest, the inside of a neuron sits at about -70 millivolts relative to the
outside, a state called the resting membrane potential.

This voltage difference exists because ion channels and pumps control the flow of charged
particles across the membrane. The sodium-potassium pump moves three sodium ions out of the
cell for every two potassium ions it moves in, using ATP as its energy source. Because the
membrane is more permeable to potassium than to sodium at rest, potassium tends to leak
outward, pulling the inside of the cell toward a more negative voltage.

When a neuron receives enough excitatory input, the membrane potential rises toward a
threshold, usually around -55 millivolts. If that threshold is crossed, voltage-gated sodium
channels open rapidly, sodium rushes into the cell, and the membrane potential spikes upward
in an event called an action potential. This is an all-or-nothing event: either the
threshold is crossed and a full spike fires, or it is not crossed and nothing happens,
regardless of how strong the input was.

After the peak of the spike, voltage-gated potassium channels open and sodium channels
inactivate, driving the membrane potential back down and briefly past the resting level, a
period called hyperpolarization. During part of this recovery, called the refractory period,
the neuron cannot fire again, which limits the maximum rate at which it can spike and gives
each action potential a clean, discrete shape as it travels down the axon.

Action potentials propagate along the axon without losing amplitude, because each patch of
membrane regenerates the spike rather than passively transmitting a fading signal. In
myelinated axons, the wrapping of myelin around the axon forces the spike to jump between
unmyelinated gaps called nodes of Ranvier, a process called saltatory conduction that speeds
up transmission considerably compared to an unmyelinated axon of the same diameter.""",
    ),
    (
        "Synaptic Transmission",
        """When an action potential reaches the end of an axon, it arrives at a structure
called the axon terminal, which sits close to another neuron's dendrite or cell body across
a narrow gap called the synaptic cleft. Most communication across this gap is chemical
rather than electrical.

The arrival of the spike opens voltage-gated calcium channels in the terminal. Calcium
flowing into the terminal triggers vesicles filled with neurotransmitter to fuse with the
presynaptic membrane and release their contents into the synaptic cleft, a process called
exocytosis. The neurotransmitter molecules diffuse across the cleft and bind to receptors on
the postsynaptic membrane.

Two broad receptor types exist. Ionotropic receptors are ligand-gated ion channels: when
neurotransmitter binds, the channel opens directly and ions flow immediately, producing a
fast postsynaptic response measured in milliseconds. Metabotropic receptors instead trigger
a second-messenger cascade inside the postsynaptic cell, producing slower but longer-lasting
effects, sometimes changing gene expression.

Whether a synapse is excitatory or inhibitory depends on which ions flow, not on the
identity of the neurotransmitter alone. Glutamate binding to an ionotropic receptor
typically opens channels that let sodium in, depolarizing the postsynaptic cell and
producing an excitatory postsynaptic potential. GABA binding to its ionotropic receptor
typically opens channels that let chloride in, hyperpolarizing the cell and producing an
inhibitory postsynaptic potential. A single postsynaptic neuron integrates thousands of
these excitatory and inhibitory inputs across its dendrites, and whether it fires depends on
the net sum reaching threshold at the axon hillock.

After release, neurotransmitter is cleared from the cleft by reuptake transporters that pump
it back into the presynaptic terminal, by enzymatic breakdown, or by diffusion away from the
synapse. This clearance step resets the synapse for the next signal and is also the target
of many psychoactive drugs, which slow or block reuptake to prolong the
neurotransmitter's effect.""",
    ),
    (
        "Hebbian Learning and Plasticity",
        """Hebbian learning is often summarized as "cells that fire together wire together."
The idea, proposed by Donald Hebb in 1949, is that when a presynaptic neuron repeatedly
contributes to firing a postsynaptic neuron, the connection between them strengthens,
making future firing of the postsynaptic neuron by that same presynaptic input more likely.

The best-studied physiological form of this idea is long-term potentiation, or LTP. In a
typical experiment, a brief high-frequency burst of stimulation to a presynaptic pathway
leaves the synapse measurably stronger for hours or longer, seen as a bigger postsynaptic
response to the same test stimulus. LTP in many pathways depends on the NMDA receptor, a
glutamate receptor that is unusual because it requires both neurotransmitter binding and
sufficient postsynaptic depolarization to open, since a magnesium ion normally blocks its
channel at resting potential.

This dual requirement makes the NMDA receptor a coincidence detector: it only lets calcium
into the postsynaptic cell when the presynaptic terminal releases glutamate at roughly the
same time as the postsynaptic neuron is already active. That calcium influx triggers
signaling cascades that insert additional AMPA receptors into the postsynaptic membrane, so
the same amount of glutamate produces a larger response on the next signal, the physical
basis of the strengthened connection.

The opposite process, long-term depression or LTD, weakens synapses when presynaptic
activity occurs without matching postsynaptic activity, or when the pattern of coincident
firing is different. Together, LTP and LTD let synaptic weights move in both directions,
which is necessary for a network to both learn new associations and forget or revise old
ones rather than only accumulating strength.

Because plasticity here is described at the level of individual synapses and depends on the
precise timing between presynaptic and postsynaptic activity, this same coincidence-
detection idea reappears in a more general form in spike-timing-dependent plasticity, and it
inspired the class of artificial learning rules that adjust connection weights based on
correlated activity.""",
    ),
]

PPTX_SLIDES = [
    (
        "What Is Machine Learning?",
        [
            "A program that improves at a task from experience (data), instead of following "
            "only hand-written rules.",
            "Three broad families: supervised, unsupervised, and reinforcement learning.",
            "Performance is measured on data the model did not train on, not on the "
            "training set.",
        ],
        "Emphasize the contrast with classical programming: in ML the data and the desired "
        "output are given, and the program (the model's parameters) is what gets produced.",
    ),
    (
        "Supervised vs. Unsupervised",
        [
            "Supervised learning: every training example has a label; regression predicts a "
            "number, classification predicts a category.",
            "Unsupervised learning: no labels; the goal is to find structure, such as "
            "clusters or a lower-dimensional representation.",
            "Semi-supervised and self-supervised methods use a small labeled set plus a "
            "large unlabeled set.",
        ],
        None,
    ),
    (
        "Gradient Descent",
        [
            "A loss function measures how wrong the model's predictions are on a batch of "
            "examples.",
            "The gradient of the loss with respect to each parameter points in the "
            "direction of steepest increase.",
            "Each step moves the parameters a small distance in the opposite direction, "
            "scaled by a learning rate.",
            "Too large a learning rate overshoots and diverges; too small converges too "
            "slowly.",
        ],
        "Stochastic gradient descent computes the gradient on a small random batch instead "
        "of the full dataset each step, which is noisier per step but far cheaper, and the "
        "noise itself can help escape shallow local minima.",
    ),
    (
        "Backpropagation",
        [
            "Chain rule applied layer by layer.",
            "Reuses intermediate values from the forward pass.",
        ],
        None,
    ),
    (
        "Overfitting and Regularization",
        [
            "A model overfits when it fits the noise in the training set and its "
            "performance on held-out data gets worse even as training loss keeps improving.",
            "Regularization techniques (L2 weight penalties, dropout, early stopping) trade "
            "a bit of training performance for better generalization.",
            "A validation set, separate from both training and test data, is used to choose "
            "these hyperparameters without leaking information from the test set.",
        ],
        "Early stopping means halting training when validation loss stops improving, even "
        "if training loss would keep going down -- the gap between the two curves is the "
        "signature of overfitting.",
    ),
]

VTT_CUES = [
    (0.0, 4.0, "Today we're covering the basics of reinforcement learning."),
    (
        4.0,
        9.0,
        "Unlike supervised learning, there's no dataset of labeled examples handed to "
        "the agent up front.",
    ),
    (
        9.0,
        14.5,
        "Instead, an agent interacts with an environment, taking actions and observing "
        "what happens.",
    ),
    (14.5, 19.0, "After each action, the environment returns a new state and a scalar reward."),
    (
        19.0,
        24.0,
        "The agent's goal is to choose actions that maximize the total reward it collects "
        "over time.",
    ),
    (
        24.0,
        29.5,
        "A policy is the agent's strategy: a mapping from states to actions, or to a "
        "distribution over actions.",
    ),
    (
        29.5,
        35.0,
        "A value function estimates how much future reward the agent can expect from a "
        "given state.",
    ),
    (
        35.0,
        40.0,
        "The action-value function, often written Q, estimates expected future reward "
        "for taking a specific action in a state.",
    ),
    (
        40.0,
        45.5,
        "One central tension in reinforcement learning is exploration versus exploitation.",
    ),
    (
        45.5,
        51.0,
        "Exploitation means choosing the action currently believed to be best, based on "
        "what's been learned so far.",
    ),
    (
        51.0,
        56.5,
        "Exploration means trying an action specifically to learn more about the "
        "environment, even if it looks worse right now.",
    ),
    (
        56.5,
        62.0,
        "An epsilon-greedy policy exploits most of the time, but explores a random action "
        "with small probability epsilon.",
    ),
    (
        62.0,
        67.5,
        "Temporal-difference learning updates a value estimate using the difference "
        "between consecutive predictions, without waiting for the episode to end.",
    ),
    (
        67.5,
        73.0,
        "Q-learning is a temporal-difference method that learns the action-value function "
        "directly, independent of the policy being followed.",
    ),
    (
        73.0,
        78.0,
        "Because it learns off-policy, Q-learning can reuse experience collected under an "
        "older or more exploratory policy.",
    ),
    (
        78.0,
        83.5,
        "A discount factor, usually written gamma and between zero and one, controls how "
        "much the agent values future reward compared to immediate reward.",
    ),
    (
        83.5,
        88.5,
        "A discount factor close to zero makes the agent short-sighted; close to one makes "
        "it weigh distant rewards almost as much as immediate ones.",
    ),
    (
        88.5,
        93.0,
        "When the state space is too large to hold a table of values, a function "
        "approximator, such as a neural network, estimates the value function instead.",
    ),
    (
        93.0,
        98.0,
        "That combination of deep neural networks with reinforcement learning is usually "
        "called deep reinforcement learning.",
    ),
]


def build_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for title, body in PDF_PAGES:
        page = doc.new_page()
        rect = pymupdf.Rect(36, 36, page.rect.width - 36, page.rect.height - 36)
        page.insert_textbox(rect, f"{title}\n\n{body}", fontsize=11, fontname="helv")
    doc.save(path)
    doc.close()


def build_pptx(path: Path) -> None:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title + content
    for title, bullets, notes in PPTX_SLIDES:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    presentation.save(path)


def _format_timestamp(seconds: float) -> str:
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = round((seconds - int(seconds)) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"


def build_vtt(path: Path) -> None:
    lines = ["WEBVTT", ""]
    for t0, t1, text in VTT_CUES:
        lines.append(f"{_format_timestamp(t0)} --> {_format_timestamp(t1)}")
        lines.append(text)
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    build_pdf(FIXTURES_DIR / "synthetic.pdf")
    build_pptx(FIXTURES_DIR / "synthetic.pptx")
    build_vtt(FIXTURES_DIR / "synthetic.vtt")
    print(f"wrote fixtures to {FIXTURES_DIR}")


if __name__ == "__main__":
    main()
